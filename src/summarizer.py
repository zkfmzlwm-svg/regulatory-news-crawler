import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import yaml

from .models import Article
from .translate import TranslationError, translate_texts
from .utils import clean_text, is_pdf_url, strip_tags

logger = logging.getLogger(__name__)

MAX_BULLETS = 5
# 마침표 없이 끝나는 줄이 이 단어 수보다 짧으면 문장이 아니라 소제목·사진 설명 등으로 본다
_HEADING_MAX_WORDS = 12

# "U.S.", "Jan. 5", "Dr. Smith" 같은 약어 뒤에서는 문장을 자르지 않는다
_ABBREVIATIONS = [
    "Mr", "Ms", "Mrs", "Dr", "St", "No", "vs", "Inc", "Ltd", "Co", "Corp", "Fig", "approx",
    "Jan", "Feb", "Mar", "Apr", "Jun", "Jul", "Aug", "Sep", "Sept", "Oct", "Nov", "Dec", "e.g", "i.e",
    "Ph", "Eur", "Pharm", "Ref", "Vol", "Art", "Sec",
    "Gov", "Sen", "Rep", "Gen", "Jr", "Sr", "Prof", "Dept", "Univ",
]

# 기사 내용이 아닌 사이트 공통 문구 (사진 출처, 보도자료 머리말, 구독 유도, gov.uk 접근성 안내 등)
_BOILERPLATE = re.compile(
    r"^(image credit|photo credit|credit:|\(ap photo|\(photo|media inquiries|"
    r"consumer inquiries|sign up to read|subscribe|already a subscriber|read more|share this|"
    r"updated monday through friday|this file may not be suitable|request an accessible format|"
    r"if you use assistive technology|please tell us what format|it will help us if you say|"
    r"related (articles|links|content)|content current as of|regulated product\(s\)|"
    r"get free access|create a free account|log ?in to read|"
    r"an official website of the|here['’]?s how you know|(the )?\.gov means it['’]?s official|"
    r"federal government websites often end in|before sharing sensitive information|"
    r"the site is secure|the https:// ensures|"
    # FDA 리콜 페이지 상단 요약표의 항목명이 값과 붙어 나오는 줄 (예: "Recall Reason DescriptionPotential ...")
    r"recall reason description|brand name\(s\)|product description|"
    # Health Canada 리콜 페이지의 공통 안내 문구
    r"verify if your product is affected|contact the (recalling firm|manufacturer)|"
    r"report any (other )?health product|receive emails about|health products - |"
    r"consult your health ?care (provider|professional) (prior to|before|if)|"
    # 사진 설명, RAPS 'Regulatory Recon' 고정 안내 문구
    r"from l-r|\(l-r\)|regulatory recon is|a story['’]s inclusion in regulatory recon|"
    # PMDA 영문 사이트 공통 안내 (자바스크립트·기계번역 안내, 영문 번역본 면책 문구)
    r"please make javascript on|please note that this site may not be accurate|"
    r"the following english translations? |in the event of inconsistency between the japanese|"
    r"pmda shall not be responsible|the review reports? (were selected|contains)|"
    # MHLW 공통 문구 (문의처·전화번호, 자바스크립트·자동번역·PDF 뷰어 안내, 보도자료 수신인)
    r"照会先|[（(](代表|直通)電話[）)]|報道関係者\s*各位|このページではjavascript|"
    r"このホームページを、英語|pdfファイルを見るためには|※?スパムメール防止のため|(?:お手数ですが|恐れ入りますが)、送信)",
    re.IGNORECASE,
)
# 보도자료 머리말: 문단 맨 앞만 떼어내고 뒤의 본문 문장은 살린다
_RELEASE_PREFIX = re.compile(r"^(?:for\s+)?immediate\s+release\s*[-–—:]*\s*", re.IGNORECASE)
# 유료 구독벽 등으로 문장 중간에서 잘린 경우 ("... operates …")
_TRUNCATED = re.compile(r"(?:…|\.\.\.)\s*$")
_LIST_MARKER = re.compile(r"^(?:[-–—•*·▪►・▶●○■◆]\s*)+")
_SENTENCE_SPLIT = re.compile(
    r"(?<!\.[A-Z]\.)(?<!\b[ap]\.m\.)"
    + "".join(rf"(?<!\b{re.escape(a)}\.)" for a in _ABBREVIATIONS)
    + r"(?<=[.!?])\s+(?=[A-Z0-9\"“‘'(\[])"
)


# 일본어(가나·한자) 본문: 띄어쓰기가 없어 단어 수 대신 글자 수로 판단하고 "。" 에서 문장을 나눈다
_CJK_CHAR = re.compile(r"[\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
_CJK_SENTENCE_SPLIT = re.compile(r"(?<=[。！？])")
_CJK_END = ("。", "！", "？")
# 마침표 없는 일본어 줄이 이보다 짧으면 부서명·일시·번호 붙은 항목명 등이라 버리고, 이보다 길면 소제목으로 둔다
_CJK_MIN_CHARS = 15
# 한 문장이 여러 줄로 끊겨 있는 경우(보도자료 PDF 를 옮긴 페이지) 이어 붙일 앞줄의 최소 길이
_CJK_WRAP_MIN_CHARS = 30
_CJK_LINE_START_MARKER = re.compile(r"^[・▶●○■◆※【（(０-９0-9]")
# 담당자 연락처 줄 (예: "予防接種課 課長補佐 幕内（8157）")
_JP_CONTACT = re.compile(r"[（(](?:内線)?\s*[0-9０-９]{4}\s*[）)]")
# 문장 중간에서 끊긴 줄은 대개 조사·어미(히라가나)나 쉼표로 끝난다. 명사로 끝나는 줄은 소제목으로 보고 잇지 않는다
_CJK_CONTINUES = re.compile(r"[\u3041-\u309f、，]$")
# 일본 연호: Google 번역이 "令和８年度" 를 "2017년" 처럼 틀리게 옮겨, 번역 전에 서기로 바꾼다
_JP_ERA = re.compile(r"(令和|平成)\s*([0-9０-９]+|元)\s*年")
_JP_ERA_BASE = {"令和": 2018, "平成": 1988}


def _to_western_year(text: str) -> str:
    return _JP_ERA.sub(lambda m: f"{_JP_ERA_BASE[m.group(1)] + (1 if m.group(2) == '元' else int(m.group(2)))}年", text)


def _is_cjk(text: str) -> bool:
    return len(_CJK_CHAR.findall(text)) * 3 >= len(text)


def _has_cjk(text: str) -> bool:
    return bool(_CJK_CHAR.search(text))


@dataclass
class SummaryFormat:
    title: str
    output_format: str


@dataclass
class SummaryEntry:
    index: int
    article: Article
    body: str

    @property
    def tag(self) -> str:
        return self.article.tag or self.article.source

    @property
    def title(self) -> str:
        """요약 제목 줄에 쓰는 원문(영문) 제목. 번역하지 않고 그대로 사용."""
        return clean_text(self.article.title)


def load_format(config_path: str) -> SummaryFormat:
    with open(config_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    return SummaryFormat(
        title=data.get("title", "규제동향 요약"),
        output_format=data.get("output_format", "md"),
    )


def _is_fragment(sentence: str) -> bool:
    """목록 항목이 쪼개져 문장 중간부터 시작하는 조각 (예: "prevention, treatment or ..."). "mRNA", "eCTD" 는 문장으로 본다."""
    return len(sentence) > 1 and sentence[0].islower() and sentence[1].islower()


def _is_title(para: str, title: str) -> bool:
    """본문에 다시 나오는 기사 제목. RSS 제목 뒤에 " - 09/17/2026" 같은 꼬리가 붙은 경우도 같은 제목으로 본다."""
    low = para.lower()
    return low == title or (len(low) > 20 and title.startswith(low))


def _join_wrapped_cjk(paras: List[str], title: str) -> List[str]:
    """일본어 문장이 중간에서 줄바꿈된 경우 ("…広報啓発を" / "推進してまいります。") 한 줄로 잇는다."""
    out: List[str] = []
    for para in paras:
        prev = out[-1] if out else ""
        if (
            len(prev) >= _CJK_WRAP_MIN_CHARS
            and _is_cjk(prev)
            and _CJK_CONTINUES.search(prev)
            and not _is_title(prev, title)
            and para
            and _is_cjk(para)
            and not _CJK_LINE_START_MARKER.match(para)
        ):
            out[-1] = prev + para
        else:
            out.append(para)
    return out


def _cjk_sentences(para: str):
    """일본어 문단을 (문장, 소제목) 으로 나눈다. 문장이 아닌 짧은 줄은 둘 다 아님."""
    sentences, headings = [], []
    for s in _CJK_SENTENCE_SPLIT.split(para):
        s = s.strip()
        if len(s) < 10:  # "「一瞬の判断。" 처럼 따옴표 안에서 끊긴 조각
            continue
        if s.endswith(_CJK_END):
            sentences.append(s)
        elif len(s) >= _CJK_MIN_CHARS:
            headings.append(s)
    return sentences, headings


def _extract_sentences(article: Article, full_text: Optional[str]) -> List[str]:
    text = strip_tags(full_text or article.excerpt or "")
    title = clean_text(article.title).lower()
    sentences: List[str] = []
    if not full_text and title:
        # 원문이 봇 차단 등으로 안 열리면(예: FiercePharma) RSS 의 제목 + 소개글로 대신 요약한다
        sentences.append(clean_text(article.title))
    seen = {s.lower() for s in sentences}
    truncated: List[str] = []
    headings: List[str] = []
    paras = [_LIST_MARKER.sub("", re.sub(r"\s+", " ", p).strip()) for p in text.splitlines()]
    if _has_cjk(text):
        paras = _join_wrapped_cjk(paras, title)
    for para in paras:
        para = _RELEASE_PREFIX.sub("", para)
        # 본문 첫 줄에 반복되는 기사 제목과, 문장이 아닌 짧은 소제목("Background", "What you should do" 등)은 건너뛴다
        if not para or _is_title(para, title) or _BOILERPLATE.match(para):
            continue
        if _is_cjk(para):
            if _JP_CONTACT.search(para):
                continue
            found, heads = _cjk_sentences(para)
            for x in found:
                if x not in seen:
                    seen.add(x)
                    sentences.append(x)
            headings.extend(heads)
            if len(sentences) >= MAX_BULLETS:
                break
            continue
        words = len(para.split())
        # 회사명("Pharmascience Inc.") 같은 두 단어 이하 줄도 문장이 아니다
        if words < 3 or (words < 6 and not para.endswith((".", "!", "?"))):
            continue
        if para.endswith(":") and words < 10:
            continue
        # 교육과정·행사 홍보 문구 (예: "With Updates on the new ICH Q1 Guideline!")
        if para.endswith("!") and words < 10:
            continue
        # 시간·날짜·번호만 있는 줄 (예: "9:00 a.m. - 2:00 p.m.")
        if sum(c.isalpha() for c in para) < len(para) / 2:
            continue
        for s in _SENTENCE_SPLIT.split(para):
            s = s.strip()
            if len(s) < 3 or _BOILERPLATE.match(s) or _is_fragment(s) or s.lower() in seen:
                continue
            seen.add(s.lower())
            end = s.rstrip("\"”’')]")
            if end.endswith(":"):  # "These standards replace the following:" 처럼 뒤 목록 없이는 뜻이 없는 문장
                continue
            if _TRUNCATED.search(s):
                truncated.append(s)
            elif not end.endswith((".", "!")) and len(s.split()) < _HEADING_MAX_WORDS:
                # 마침표 없는 짧은 줄(소제목·사진 설명·주소·행사 일시)과 짧은 질문형 소제목("What are listed medicines?")
                headings.append(s)
            else:
                sentences.append(s)
        if len(sentences) >= MAX_BULLETS:
            break
    # 소제목(헤드라인 모음 기사 등)이나 잘린 문장은 제대로 된 문장이 하나도 없을 때만 쓴다
    result = (sentences or headings or truncated)[:MAX_BULLETS]
    if not result and full_text:
        # 본문이 PDF 링크 목록뿐인 페이지(PMDA·MHLW 자료 게시 페이지 등)는 제목 + 소개글로 대신한다
        return _extract_sentences(article, None)
    return result


def _summarize_fallback(article: Article, full_text: Optional[str]) -> str:
    """API 키가 없을 때 쓰는 가장 단순한 추출 요약 (번역 없이 원문 문장 그대로)."""
    bullets = _extract_sentences(article, full_text)
    return "\n".join(f"- {s}" for s in bullets) if bullets else "- (본문을 가져오지 못했습니다)"


def _summarize_free_translate(article: Article, full_text: Optional[str]) -> str:
    """API 키 없이 무료 번역(비공식 Google 번역 등)으로 원문 문장을 한국어로 옮긴 요약.

    실제 AI 요약처럼 내용을 재구성/분석하지는 않고 추출한 문장을 그대로 번역만 하므로
    품질은 AI 요약보다 단순하지만, 번역된 한국어 결과를 얻는 데 비용이 들지 않는다.
    """
    bullets = _extract_sentences(article, full_text)
    if not bullets:
        return "- (본문을 가져오지 못했습니다)"

    # 일본어 제목(MHLW 등)은 요약 제목 줄에 원문 그대로 나가므로, 번역한 제목을 첫 줄에 붙인다
    title = clean_text(article.title)
    with_title = _has_cjk(title) and title not in bullets
    texts = [_to_western_year(t) for t in ([title] if with_title else []) + bullets]
    try:
        translated = translate_texts(texts, target="ko")
    except TranslationError as exc:
        logger.warning("번역 실패, 원문 유지 (%s): %s", article.title, exc)
        return "\n".join(f"- {s} (번역 실패)" for s in bullets)
    lines = [f"- 제목: {translated[0] or title}"] if with_title else []
    translated = translated[1:] if with_title else translated
    lines.extend(f"- {t or s}" for s, t in zip(bullets, translated))
    return "\n".join(lines)


def summarize_article(article: Article, full_text: Optional[str], mode: str = "free") -> str:
    """체크한 기사 본문을 요약해 body(불릿 텍스트)를 반환. API 키 불필요.

    mode:
      - "simple": 번역 없이 원문 문장을 그대로 추출
      - "free": Google 번역(비공식, 무료)으로 문장을 한국어로 번역
    """
    if not full_text and not article.excerpt and is_pdf_url(article.url):
        # 제목 한 줄만 번역해 보여주는 것보다 원문을 직접 열어보라고 안내 (예: NMPA 뉴스레터)
        return "- (PDF 문서라 본문을 자동으로 요약할 수 없습니다. 링크에서 직접 확인하세요)"
    if mode == "simple":
        return _summarize_fallback(article, full_text)
    return _summarize_free_translate(article, full_text)


def render_markdown(fmt: SummaryFormat, entries: List[SummaryEntry]) -> str:
    lines = [f"# {fmt.title}", ""]
    for e in entries:
        lines.append(f"{e.index}. [{e.tag}] {e.title}")
        if e.body:
            lines.append(e.body)
        lines.append(f"- 링크: {e.article.url}")
        lines.append("")
    return "\n".join(lines)


def render_txt(fmt: SummaryFormat, entries: List[SummaryEntry]) -> str:
    lines = [fmt.title, "=" * len(fmt.title), ""]
    for e in entries:
        lines.append(f"{e.index}. [{e.tag}] {e.title}")
        if e.body:
            lines.append(e.body)
        lines.append(f"- 링크: {e.article.url}")
        lines.append("")
    return "\n".join(lines)


def _add_hyperlink(paragraph, url: str, text: str):
    """python-docx 는 하이퍼링크를 기본 지원하지 않아 저수준 XML로 추가."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(underline)
    run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def render_docx(fmt: SummaryFormat, entries: List[SummaryEntry], output_path: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for style_name in ("List Bullet", "List Bullet 2"):
        if style_name in doc.styles:
            doc.styles[style_name].font.size = Pt(11)

    doc.add_heading(fmt.title, level=1)
    for e in entries:
        doc.add_heading(f"{e.index}. [{e.tag}] {e.title}", level=2)
        for line in e.body.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith(": "):
                doc.add_paragraph(stripped[2:], style="List Bullet 2")
            elif stripped.startswith("- "):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            else:
                doc.add_paragraph(stripped, style="List Bullet")
        link_p = doc.add_paragraph()
        link_p.add_run("링크: ")
        try:
            _add_hyperlink(link_p, e.article.url, e.article.url)
        except Exception:
            link_p.add_run(e.article.url)
    doc.save(output_path)


def output_extension(fmt: SummaryFormat) -> str:
    return {"docx": "docx", "txt": "txt"}.get(fmt.output_format.lower(), "md")


def write_summaries(fmt: SummaryFormat, entries: List[SummaryEntry], output_path: str) -> str:
    output_format = fmt.output_format.lower()
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    if output_format == "docx":
        render_docx(fmt, entries, output_path)
    elif output_format == "txt":
        Path(output_path).write_text(render_txt(fmt, entries), encoding="utf-8")
    else:
        Path(output_path).write_text(render_markdown(fmt, entries), encoding="utf-8")
    return output_path
